from django.shortcuts import render

# Create your views here.
"""
App: rapports
==============
Génération de rapports téléchargeables : fiche véhicule en PDF (ReportLab)
et export de toute la flotte en Excel (openpyxl).
"""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.exceptions import NotFound

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from vehicules.models import Vehicule


def _formater_fcfa(montant) -> str:
    return f"{montant:,.0f}".replace(',', ' ')


def _style_tableau_historique() -> TableStyle:
    return TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976d2')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e0e0e0')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
    ])


class VehiculeExportPDFView(APIView):
    """GET /api/rapports/vehicules/<uuid>/export-pdf/ — fiche complète en PDF."""
    permission_classes = [IsAuthenticated]

    def get(self, request, vehicule_id):
        try:
            vehicule = Vehicule.objects.get(id=vehicule_id, is_active=True)
        except Vehicule.DoesNotExist:
            raise NotFound("Véhicule introuvable.")

        maintenances = list(vehicule.maintenances.filter(is_active=True).order_by('-date_intervention'))
        pleins = list(vehicule.pleins_carburant.filter(is_active=True).order_by('-date_plein'))
        documents = list(vehicule.documents.filter(is_active=True).order_by('date_expiration'))

        buffer = BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('TitleCustom', parent=styles['Title'], fontSize=18, spaceAfter=6)
        subtitle_style = ParagraphStyle(
            'SubtitleCustom', parent=styles['Normal'], fontSize=11, textColor=colors.grey
        )
        section_style = ParagraphStyle('SectionCustom', parent=styles['Heading2'], spaceBefore=16, spaceAfter=8)

        elements = []
        elements.append(Paragraph(f"Fiche véhicule — {vehicule.immatriculation}", title_style))
        elements.append(Paragraph(
            f"{vehicule.marque} {vehicule.modele} ({vehicule.annee}) — "
            f"Généré le {timezone.now().strftime('%d/%m/%Y à %H:%M')}",
            subtitle_style,
        ))
        elements.append(Spacer(1, 0.5 * cm))

        # --- Informations générales ---
        infos_data = [
            ['Statut', vehicule.get_statut_display()],
            ['Type', vehicule.get_type_vehicule_display()],
            ['Carburant', vehicule.get_type_carburant_display()],
            ['Kilométrage actuel', f"{vehicule.kilometrage_actuel:,} km".replace(',', ' ')],
            ["Site d'affectation", vehicule.site_affectation or '—'],
            ['N° châssis', vehicule.numero_chassis],
            ["Date d'acquisition", vehicule.date_acquisition.strftime('%d/%m/%Y')],
            ["Prix d'acquisition", f"{_formater_fcfa(vehicule.prix_acquisition)} FCFA"],
        ]
        infos_table = Table(infos_data, colWidths=[5 * cm, 10 * cm])
        infos_table.setStyle(TableStyle([
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('LINEBELOW', (0, 0), (-1, -1), 0.5, colors.HexColor('#e0e0e0')),
        ]))
        elements.append(infos_table)

        # --- Maintenance ---
        cout_maintenance_total = sum(m.cout for m in maintenances)
        elements.append(Paragraph(
            f"Historique de maintenance — Total : {_formater_fcfa(cout_maintenance_total)} FCFA",
            section_style,
        ))
        if maintenances:
            data = [['Date', 'Type', 'Description', 'Coût (FCFA)']]
            for m in maintenances:
                data.append([
                    m.date_intervention.strftime('%d/%m/%Y'),
                    m.get_type_maintenance_display(),
                    Paragraph(m.description, styles['Normal']),
                    _formater_fcfa(m.cout),
                ])
            table = Table(data, colWidths=[2.2 * cm, 3 * cm, 7 * cm, 2.8 * cm], repeatRows=1)
            table.setStyle(_style_tableau_historique())
            elements.append(table)
        else:
            elements.append(Paragraph("Aucune intervention enregistrée.", styles['Normal']))

        # --- Carburant ---
        cout_carburant_total = sum(p.cout_total for p in pleins)
        elements.append(Paragraph(
            f"Historique carburant — Total : {_formater_fcfa(cout_carburant_total)} FCFA",
            section_style,
        ))
        if pleins:
            data = [['Date', 'Litres', 'Coût (FCFA)', 'Station']]
            for p in pleins:
                data.append([
                    p.date_plein.strftime('%d/%m/%Y'),
                    f"{p.litres} L",
                    _formater_fcfa(p.cout_total),
                    p.station or '—',
                ])
            table = Table(data, colWidths=[2.5 * cm, 2.5 * cm, 3 * cm, 7 * cm], repeatRows=1)
            table.setStyle(_style_tableau_historique())
            elements.append(table)
        else:
            elements.append(Paragraph("Aucun plein enregistré.", styles['Normal']))

        # --- Documents ---
        elements.append(Paragraph("Documents", section_style))
        if documents:
            data = [['Type', 'Numéro', 'Expiration', 'Statut']]
            for d in documents:
                data.append([
                    d.get_type_document_display(),
                    d.numero_document or '—',
                    d.date_expiration.strftime('%d/%m/%Y') if d.date_expiration else '—',
                    'Valide' if d.est_valide else 'Expiré',
                ])
            table = Table(data, colWidths=[4 * cm, 4 * cm, 3 * cm, 4 * cm], repeatRows=1)
            table.setStyle(_style_tableau_historique())
            elements.append(table)
        else:
            elements.append(Paragraph("Aucun document enregistré.", styles['Normal']))

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer, content_type='application/pdf')
        filename = f"fiche_{vehicule.immatriculation.replace(' ', '_')}.pdf"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response


class FlotteExportExcelView(APIView):
    """GET /api/rapports/flotte/export-excel/ — export Excel de toute la flotte."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        vehicules = Vehicule.objects.filter(is_active=True).order_by('immatriculation')

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Flotte"

        headers = [
            'Immatriculation', 'Marque', 'Modèle', 'Année', 'Statut',
            'Kilométrage', 'Site', 'Coût maintenance (FCFA)',
            'Coût carburant (FCFA)', 'Coût total (FCFA)',
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color='1976D2', end_color='1976D2', fill_type='solid')
        header_font = Font(color='FFFFFF', bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        total_maintenance_flotte = 0
        total_carburant_flotte = 0

        for v in vehicules:
            cout_maintenance = sum(m.cout for m in v.maintenances.filter(is_active=True))
            cout_carburant = sum(p.cout_total for p in v.pleins_carburant.filter(is_active=True))
            total_maintenance_flotte += cout_maintenance
            total_carburant_flotte += cout_carburant

            ws.append([
                v.immatriculation, v.marque, v.modele, v.annee,
                v.get_statut_display(), v.kilometrage_actuel,
                v.site_affectation or '—',
                float(cout_maintenance), float(cout_carburant),
                float(cout_maintenance + cout_carburant),
            ])

        ws.append([])
        ws.append([
            'TOTAL FLOTTE', '', '', '', '', '', '',
            float(total_maintenance_flotte), float(total_carburant_flotte),
            float(total_maintenance_flotte + total_carburant_flotte),
        ])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

        for i, header in enumerate(headers, start=1):
            ws.column_dimensions[chr(64 + i)].width = max(len(header) + 2, 16)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        filename = f"export_flotte_{timezone.now().strftime('%Y%m%d')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response